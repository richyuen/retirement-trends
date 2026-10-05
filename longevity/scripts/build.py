"""Build longevity tables from NCHS life tables and SSA Trustees Reports.

Inputs (raw/, see README for sources):
  hus2020-21_LExpMort.xlsx           Health, United States 2020-2021, Table LExpMort (1900-2019, 1 decimal)
  nvsr/<vol-no>/T01..T18.xlsx        NCHS United States Life Tables, data years 2018-2024
  trustees/tr20{20,26}_tables_VA4_VA5.txt  Tables V.A4 (period) and V.A5 (cohort) text from the
                                     2020 and 2026 OASDI Trustees Reports (govinfo House Documents)
Outputs (output/):
  le_nchs_1980_2024.csv              period life expectancy at birth, 65, 75 by sex, annual
  life_table_summary_2018_2024.csv   e0/e65/e75 and survival 65->85/90/95 by group, sex, year
  trustees_period_cohort.csv         SSA period + cohort life expectancy, 2020 and 2026 reports
  covid_gap.csv                      actual 2020-2024 vs pre-COVID trend and 2020 Trustees path
  income_le40_chetty.csv             life expectancy at 40 by household income quartile, 2001-2014
"""
import csv, os, re
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
RAW = os.path.join(HERE, '..', 'raw')
OUT = os.path.join(HERE, '..', 'output')

NVSR = {2018: '69-12', 2019: '70-19', 2020: '71-01', 2021: '72-12',
        2022: '74-02', 2023: '74-06', 2024: '75-05'}
# Table numbers by (group, sex). 2018 report used a different race grouping, so only total is used for 2018.
GROUPS = {'All': 1, 'Hispanic': 4, 'NH AIAN': 7, 'NH Asian': 10, 'NH Black': 13, 'NH White': 16}
SEXES = {'Both': 0, 'Male': 1, 'Female': 2}


def write(name, rows, cols):
    with open(os.path.join(OUT, name), 'w', newline='') as f:
        w = csv.DictWriter(f, cols)
        w.writeheader()
        w.writerows(rows)


def life_table(path):
    rows = list(openpyxl.load_workbook(path, read_only=True).active.iter_rows(values_only=True))
    lx, ex = {}, {}
    for r in rows:
        m = re.match(r'^(\d+)\s*[–‒—-]', str(r[0]).strip()) if r[0] is not None else None
        if m and isinstance(r[2], (int, float)):
            a = int(m.group(1))
            lx[a], ex[a] = r[2], r[6]
    return lx, ex


# 1. NVSR life tables 2018-2024
summary = []
for year, vol in NVSR.items():
    for g, base in GROUPS.items():
        if year == 2018 and g != 'All':
            continue
        for s, off in SEXES.items():
            lx, ex = life_table(os.path.join(RAW, 'nvsr', vol, 'T%02d.xlsx' % (base + off)))
            summary.append(dict(year=year, group=g, sex=s,
                                e0=round(ex[0], 2), e65=round(ex[65], 2), e75=round(ex[75], 2),
                                surv65_85=round(100 * lx[85] / lx[65], 1),
                                surv65_90=round(100 * lx[90] / lx[65], 1),
                                surv65_95=round(100 * lx[95] / lx[65], 1)))
write('life_table_summary_2018_2024.csv', summary,
      ['year', 'group', 'sex', 'e0', 'e65', 'e75', 'surv65_85', 'surv65_90', 'surv65_95'])

# 2. Long annual series: HUS LExpMort 1980-2017 (all races), NVSR 2018-2024
ws = openpyxl.load_workbook(os.path.join(RAW, 'hus2020-21_LExpMort.xlsx'), read_only=True).active
hus, block = {}, None
for r in ws.iter_rows(values_only=True):
    lab = str(r[0]).strip() if r[0] is not None else ''
    if lab.startswith('Specified age') and hus:
        break  # later panels are by race/Hispanic origin only
    if lab.startswith('At birth'):
        block = 'e0'
    elif lab.startswith('At 65'):
        block = 'e65'
    elif lab.startswith('At 75'):
        block = 'e75'
    elif lab.startswith('At ') or lab.startswith('Death rates') or lab.startswith('All causes'):
        block = None
    m = re.match(r'^(\d{4})', lab)
    if block and m and isinstance(r[1], (int, float)):
        for s, i in SEXES.items():
            hus[(int(m.group(1)), s, block)] = r[1 + i]
long_rows = []
for y in range(1980, 2025):
    for s in SEXES:
        if y < 2018:
            d = dict(year=y, sex=s, e0=hus[(y, s, 'e0')], e65=hus[(y, s, 'e65')],
                     e75=hus[(y, s, 'e75')], source='NCHS Health US 2020-21 Table LExpMort')
        else:
            t = next(x for x in summary if x['year'] == y and x['group'] == 'All' and x['sex'] == s)
            d = dict(year=y, sex=s, e0=t['e0'], e65=t['e65'], e75=t['e75'],
                     source='NCHS US Life Tables %d (NVSR %s)' % (y, NVSR[y]))
        long_rows.append(d)
write('le_nchs_1980_2024.csv', long_rows, ['year', 'sex', 'e0', 'e65', 'e75', 'source'])

# Consistency check: HUS (rounded) vs NVSR for 2018-2019
for y in (2018, 2019):
    for s in SEXES:
        for k in ('e0', 'e65', 'e75'):
            nv = next(x for x in summary if x['year'] == y and x['group'] == 'All' and x['sex'] == s)[k]
            assert abs(hus[(y, s, k)] - nv) <= 0.06, (y, s, k, hus[(y, s, k)], nv)

# 3. Trustees Reports: period (V.A4) and cohort (V.A5), intermediate
tr_rows = []
for rep in (2020, 2026):
    txt = open(os.path.join(RAW, 'trustees', 'tr%d_tables_VA4_VA5.txt' % rep)).read()
    i4, i5 = txt.index('Table V.A4'), txt.index('Table V.A5')
    for kind, part in (('period', txt[i4:i5]), ('cohort', txt[i5:])):
        for line in part.splitlines():
            m = re.match(r'^\s*(\d{4})\s*[a-d]?\s*[ .]+\s+([\d. ]+)$', line)
            if not m:
                continue
            nums = [float(x) for x in m.group(2).split()]
            y = int(m.group(1))
            tr_rows.append(dict(report=rep, kind=kind, year=y, e0_male=nums[0], e0_female=nums[1],
                                e65_male=nums[2], e65_female=nums[3]))
write('trustees_period_cohort.csv', tr_rows,
      ['report', 'kind', 'year', 'e0_male', 'e0_female', 'e65_male', 'e65_female'])


def tr(rep, kind, y, col):
    hit = [r[col] for r in tr_rows if r['report'] == rep and r['kind'] == kind and r['year'] == y]
    if hit:
        return hit[0]
    ys = sorted(r['year'] for r in tr_rows if r['report'] == rep and r['kind'] == kind)
    lo = max(v for v in ys if v < y); hi = min(v for v in ys if v > y)
    a, b = tr(rep, kind, lo, col), tr(rep, kind, hi, col)
    return a + (b - a) * (y - lo) / (hi - lo)


# 4. COVID gap: actual NCHS 2020-2024 vs linear trend fitted 2010-2019, and vs 2020 Trustees (period, intermediate)
def fit(xs, ys):
    n = len(xs); mx = sum(xs) / n; my = sum(ys) / n
    b = sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / sum((x - mx) ** 2 for x in xs)
    return my - b * mx, b


gap = []
for s in SEXES:
    for k in ('e0', 'e65'):
        pts = [(r['year'], r[k]) for r in long_rows if r['sex'] == s and 2010 <= r['year'] <= 2019]
        a, b = fit(*zip(*pts))
        pts00 = [(r['year'], r[k]) for r in long_rows if r['sex'] == s and 2000 <= r['year'] <= 2010]
        _, b00 = fit(*zip(*pts00))
        for y in range(2019, 2025):
            act = next(r[k] for r in long_rows if r['sex'] == s and r['year'] == y)
            row = dict(sex=s, measure=k, year=y, actual=act, trend_2010_19=round(a + b * y, 2),
                       gap_vs_trend=round(act - (a + b * y), 2),
                       slope_2010_19_per_decade=round(10 * b, 2),
                       slope_2000_10_per_decade=round(10 * b00, 2))
            if s != 'Both':
                col = '%s_%s' % (k, s.lower())
                p = tr(2020, 'period', y, col)
                row['tr2020_projection'] = round(p, 2)
                row['gap_vs_tr2020'] = round(act - p, 2)
            gap.append(row)
write('covid_gap.csv', gap, ['sex', 'measure', 'year', 'actual', 'trend_2010_19', 'gap_vs_trend',
                             'slope_2000_10_per_decade', 'slope_2010_19_per_decade',
                             'tr2020_projection', 'gap_vs_tr2020'])
print('ok', len(summary), len(long_rows), len(tr_rows), len(gap))

# 5. Life expectancy at 40 by household income (Chetty et al. 2016, Health Inequality Project tables 1-2,
#    race-adjusted, expected age at death for 40-year-olds, 2001-2014). Quartile = mean of percentiles.
INC = os.path.join(RAW, 'income')
t1 = list(csv.DictReader(open(os.path.join(INC, 'health_ineq_online_table_1.csv'))))
t2 = list(csv.DictReader(open(os.path.join(INC, 'health_ineq_online_table_2.csv'))))
inc_rows = []
for g, sex in (('M', 'Male'), ('F', 'Female')):
    le = {int(r['pctile']): float(r['le_raceadj']) for r in t1 if r['gnd'] == g}
    for q in range(1, 5):
        ps = range(25 * q - 24, 25 * q + 1)
        by_year = {}
        for r in t2:
            if r['gnd'] == g and int(r['pctile']) in ps:
                by_year.setdefault(int(r['year']), []).append(float(r['le_raceadj']))
        yrs = sorted(by_year)
        vals = [sum(by_year[y]) / len(by_year[y]) for y in yrs]
        _, slope = fit(yrs, vals)
        inc_rows.append(dict(sex=sex, income_group='Q%d' % q,
                             le40_pooled=round(sum(le[p] for p in ps) / 25, 1),
                             le40_2001=round(vals[0], 1), le40_2014=round(vals[-1], 1),
                             gain_per_year_2001_14=round(slope, 3)))
    inc_rows.append(dict(sex=sex, income_group='Bottom 1%', le40_pooled=round(le[1], 1)))
    inc_rows.append(dict(sex=sex, income_group='Top 1%', le40_pooled=round(le[100], 1)))
write('income_le40_chetty.csv', inc_rows, ['sex', 'income_group', 'le40_pooled', 'le40_2001',
                                           'le40_2014', 'gain_per_year_2001_14'])
print('income rows', len(inc_rows))
